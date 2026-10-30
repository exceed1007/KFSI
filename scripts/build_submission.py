# -*- coding: utf-8 -*-
"""生成投稿 Computers & Security 所需配套文件。

产出（写入 submission/ 目录）：
    Cover_letter.docx     投稿信
    Title_page.docx       标题页（作者信息为占位符）
    Highlights.docx       Highlights（每条 <= 85 字符）
    Declarations.docx     利益冲突 / CRediT / 数据可用性 / 基金 / 伦理 / AI 声明
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = (Path('E:/3-学术工作/论文/正在构思的论文/'
                    '基于关键字段识别与辅助的网络协议字段边界推断方法'
                    '/submission'))

TITLE = ('Key-field identification and recursive splitting for field '
         'boundary inference of binary network protocols')

REPO = 'https://github.com/exceed1007/KFSI'

HIGHLIGHTS = [
    'Key fields are used as splitting operators rather than as an output '
    'target.',
    'Horizontal and vertical splitting remove heterogeneity and offset drift.',
    'F1 of 0.7539 on eight protocols, +56.0% over the strongest baseline.',
    'Calibration, development and test sets are disjoint at message level.',
    'The method needs no training: prior weights match the fitted weights.',
]


# ----------------------------------------------------------------------
def _doc():
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
    if rf is not None:
        rf.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/'
               'main}eastAsia', 'Times New Roman')
    for s in d.sections:
        s.left_margin = s.right_margin = Pt(72)
        s.top_margin = s.bottom_margin = Pt(72)
    return d


def _para(d, text='', *, size=12, bold=False, align=None, space_after=6,
          italic=False):
    p = d.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.name = 'Times New Roman'
    return p


def _placeholder(d, text, size=12):
    """占位符：以 [ ] 标记，提示用户补充。"""
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.name = 'Times New Roman'
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    return p


# ----------------------------------------------------------------------
def cover_letter(path):
    d = _doc()
    _para(d, '[ Corresponding author name ]', bold=True)
    _placeholder(d, '[ Affiliation, full postal address ]')
    _placeholder(d, '[ Email address ]  |  [ Telephone ]')
    _para(d)
    _para(d, '[ Date ]')
    _para(d)
    _para(d, 'Editor-in-Chief, Computers & Security')
    _para(d, 'Elsevier')
    _para(d)
    _para(d, 'Subject: Submission of a research article for consideration '
             'in Computers & Security', bold=True)
    _para(d)
    _para(d, 'Dear Editor-in-Chief,')
    _para(d)
    _para(d,
          'We submit the manuscript entitled "{}" for consideration as a '
          'research article in Computers & Security. The paper addresses a '
          'problem that sits directly in the scope of the journal: how to '
          'recover field boundaries of binary protocols for which no '
          'specification, no annotation and no implementation are available, '
          'which is precisely the situation practitioners face when '
          'analysing industrial control, IoT and proprietary '
          'protocols.'.format(TITLE))
    _para(d,
          'Field boundaries are a prerequisite for three security tasks that '
          'the journal regularly publishes on. A fuzzer that mutates a '
          'length field or a checksum produces inputs that a parser discards '
          'before reaching any deep logic; an intrusion detection system '
          'anchors its signatures and anomaly models on field offsets; and '
          'evasion or fingerprinting work needs to know which bytes a '
          'receiver actually validates. Prior work has shown that this '
          'information increases fuzzing coverage and leads to '
          'vulnerabilities that affect physical control processes, but for '
          'closed protocols it has to be obtained automatically.')
    _para(d,
          'The manuscript makes three contributions. First, it gives '
          'computable criteria for the two key fields that real protocols '
          'almost always define, the structure identifier field and the '
          'length field, turning each protocol design constraint into one '
          'testable quantity. Second, it uses these fields as splitting '
          'operators rather than as an output target: messages are split '
          'horizontally by structure and vertically by length, recursively '
          'reducing the problem to homogeneous, offset-consistent regions, '
          'after which position-level evidence is fused and decoded by '
          'maximum a posteriori inference with a field-width prior. Third, '
          'it evaluates the method against three recent baselines '
          '(BinaryInferno, REInPro, PFT) under a partition into calibration, '
          'development and test sets that are disjoint at both file and '
          'message level, with message-level decontamination, and reports '
          'ablation, sensitivity and reference-granularity analyses. KFSI '
          'attains a position F1 of 0.7539 and an interval F1 of 0.7180 on '
          '1110 messages from eight protocols, improving the strongest '
          'reproducible baseline by 56.0% (45.7% against the best oracle '
          'tuning of that baseline), and it remains the best method on five '
          'of the eight protocols.')
    _para(d,
          'We believe the paper fits Computers & Security for two further '
          'reasons. It is reproducible: the implementation, the partition '
          'manifest with source, MD5 and message counts, the '
          'decontamination script and every experiment script are public at '
          '{}, and the method reproduces without any training. It is also '
          'honest about its limits: two negative results are reported, the '
          'learned position-level weights do not beat the prior weights, and '
          'bit-level structure identifier fields are neutral on this data '
          'set; and the three protocols on which a baseline wins are '
          'analysed as differences in granularity orientation rather than as '
          'localisation errors.'.format(REPO))
    _para(d,
          'We confirm that this work is original, that it has not been '
          'published previously, that it is not under consideration by any '
          'other journal, and that all authors have read and approved the '
          'manuscript and agree to its submission.')
    _para(d)
    _placeholder(d, '[ Optional: suggested reviewers — name, affiliation, '
                    'email, and why ]')
    _placeholder(d, '[ Optional: reviewers to exclude — name, affiliation, '
                    'and reason ]')
    _para(d)
    _para(d, 'Thank you for your consideration.')
    _para(d)
    _para(d, 'Sincerely,')
    _para(d)
    _placeholder(d, '[ Corresponding author name ]')
    _placeholder(d, '[ Signature ]')
    d.save(path)
    return path


def title_page(path):
    d = _doc()
    _para(d, TITLE, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
          space_after=18)
    _para(d, 'Authors', bold=True)
    _placeholder(d, '[ Author 1: Given Name Family Name ]')
    _placeholder(d, '[ Author 2: Given Name Family Name ]')
    _placeholder(d, '[ Add or remove author lines as needed ]')
    _para(d)
    _para(d, 'Affiliations', bold=True)
    _placeholder(d, '[ a) Department, Institution, City, Country ]')
    _placeholder(d, '[ b) Department, Institution, City, Country ]')
    _para(d)
    _para(d, 'Corresponding author', bold=True)
    _placeholder(d, '[ Name, email address, telephone, ORCID iD ]')
    _para(d)
    _para(d, 'ORCID identifiers', bold=True)
    _placeholder(d, '[ Author 1: 0000-0000-0000-0000 ]')
    _placeholder(d, '[ Author 2: 0000-0000-0000-0000 ]')
    _para(d)
    _para(d, 'Present or permanent address', bold=True)
    _placeholder(d, '[ Only if different from the affiliation above; '
                    'otherwise delete this section ]')
    _para(d)
    _para(d, 'Abstract', bold=True)
    _para(d, 'The abstract is included in the manuscript file; if the '
             'system requires it on the title page as well, paste it here.')
    _para(d)
    _para(d, 'Keywords', bold=True)
    _para(d, 'protocol reverse engineering; field boundary inference; key '
             'field identification; recursive splitting; fuzzing; maximum a '
             'posteriori decoding')
    d.save(path)
    return path


def highlights(path):
    d = _doc()
    _para(d, 'Highlights', size=14, bold=True, space_after=10)
    for h in HIGHLIGHTS:
        p = d.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(h)
        r.font.size = Pt(12)
        r.font.name = 'Times New Roman'
    _para(d)
    _para(d, 'Note for the author: each highlight must not exceed 85 '
             'characters including spaces. Current lengths: '
             + ', '.join(str(len(h)) for h in HIGHLIGHTS) + '.',
          size=10, italic=True)
    d.save(path)
    return path


def declarations(path):
    d = _doc()
    _para(d, 'Declarations', size=14, bold=True, space_after=10)

    _para(d, 'Declaration of competing interests', bold=True)
    _placeholder(d, '[ Choose one and delete the other: ]')
    _para(d, 'The authors declare that they have no known competing '
             'financial interests or personal relationships that could have '
             'appeared to influence the work reported in this paper.')
    _para(d)
    _para(d, 'CRediT authorship contribution statement', bold=True)
    _placeholder(d, '[ Author 1: Conceptualization, Methodology, Software, '
                    'Writing - original draft ]')
    _placeholder(d, '[ Author 2: Validation, Data curation, Writing - '
                    'review & editing ]')
    _placeholder(d, '[ Adjust roles per author: Conceptualization, Data '
                    'curation, Formal analysis, Funding acquisition, '
                    'Investigation, Methodology, Project administration, '
                    'Resources, Software, Supervision, Validation, '
                    'Visualization, Writing - original draft, Writing - '
                    'review & editing ]')
    _para(d)
    _para(d, 'Data availability', bold=True)
    _para(d, 'The implementation, the three-way partition manifest (source, '
             'MD5 and message count of every capture), the '
             'decontamination script and all experiment scripts are '
             'publicly available at {}. The test corpora are derived from '
             'publicly distributed capture files; the manifest records the '
             'origin of each capture so that the exact evaluation data can '
             'be reconstructed. No new data were generated that require '
             'access control.'.format(REPO))
    _para(d)
    _para(d, 'Funding', bold=True)
    _placeholder(d, '[ e.g. This work was supported by the National Natural '
                    'Science Foundation of China under Grant No. XXXXXXXX. '
                    'Delete this section if no funding was received. ]')
    _para(d)
    _para(d, 'Ethics statement', bold=True)
    _para(d, 'This study used only publicly distributed network capture '
             'files of industrial control and IoT protocols. It involved no '
             'human subjects, no animals and no personal data, so ethical '
             'approval was not required.')
    _para(d)
    _para(d, 'Declaration of generative AI and AI-assisted technologies in '
             'the writing process', bold=True)
    _placeholder(d, '[ Required by Elsevier. Choose the wording that matches '
                    'what actually happened, for example: ]')
    _para(d, 'During the preparation of this work the authors used '
             'AI-assisted tools for language editing and for typesetting '
             'the manuscript. After using these tools, the authors reviewed '
             'and edited the content as needed and take full responsibility '
             'for the content of the published article.')
    _para(d)
    _para(d, 'Acknowledgements', bold=True)
    _placeholder(d, '[ Optional: colleagues, reviewers, technical staff. '
                    'Delete if not applicable. ]')
    d.save(path)
    return path


def checklist(path):
    txt = f"""# 投稿 Computers & Security — 需你补充的内容清单

生成的文件位于本目录：

| 文件 | 用途 | 是否需要补充 |
|---|---|---|
| `../KFSI_manuscript.docx` | 英文正文（含摘要、关键词、图 1–6、表 1–8、公式 1–18、参考文献 33 条） | 需检查标题译法 |
| `Cover_letter.docx` | 投稿信 | 作者/单位/邮箱/电话、日期、可选推荐与回避审稿人 |
| `Title_page.docx` | 标题页 | 全部作者姓名、单位、通讯作者、ORCID |
| `Highlights.docx` | Highlights（5 条，均 ≤85 字符） | 可按需改写 |
| `Declarations.docx` | 利益冲突、CRediT、数据可用性、基金、伦理、AI 声明 | 基金号、CRediT 分工、AI 声明措辞 |

## 必须补充（红色占位符 `[ ... ]`）

1. **作者信息**：姓名（Given Name Family Name）、单位（含院系、城市、国家）、
   通讯作者邮箱与电话、ORCID iD。封面信与标题页中出现的 `[ ... ]` 全部需要替换。
2. **基金信息**：若有资助，请给出资助机构与批准号（Declarations 的 Funding 段）。
3. **CRediT 分工**：每位作者对应的角色。
4. **生成式 AI 声明**：Elsevier 强制要求。请按实际情况确认或改写
   Declarations 中的措辞（若完全没有使用 AI 工具，应改为
   "The authors did not use generative AI or AI-assisted technologies ..."）。
5. **利益冲突**：默认写的是"无冲突"，如不适用请修改。

## 需要你确认的两个选择

1. **英文标题**：当前译作
   "{TITLE}"。
   若希望在标题中直接体现安全价值（Computers & Security 定位），可改为：
   "Field boundary inference for binary protocol security analysis: a "
   "key-field-driven recursive splitting approach"。
2. **Graphical Abstract**：Elsevier 建议提供（非强制）。候选可用稿件中的
   图 1（KFSI 总体框架）。如需我另外绘制专门的 graphical abstract，请告知。

## 已自动处理、无需你操作的部分

* 正文数字全部由 `outputs/*.json` 自动填充，与中文稿一致；
  参考文献 33 条按正文首次引用顺序自动编号。
* 数据可用性已写明仓库地址 {REPO}。
* 结论末尾"Data availability"段保留（中文稿该段被删除过，此处按期刊要求保留）。

## 提交时的文件清单（Editorial Manager）

1. Manuscript（正文，建议同时提供可编辑 docx 与 PDF）
2. Title page
3. Cover letter
4. Highlights
5. Graphical Abstract（可选图片）
6. Declarations（如需单独上传）
7. Supplementary material：本研究的实验脚本与数据已在 GitHub 公开，
   可选择上传代码仓库的归档包（zip）作为补充材料。
"""
    path.write_text(txt, encoding='utf-8')
    return path


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default=str(DEFAULT_OUT))
    args = ap.parse_args(argv)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    made = [cover_letter(out / 'Cover_letter.docx'),
            title_page(out / 'Title_page.docx'),
            highlights(out / 'Highlights.docx'),
            declarations(out / 'Declarations.docx'),
            checklist(out / 'SUBMISSION_CHECKLIST.md')]
    for m in made:
        print('saved:', m)
    return made


if __name__ == '__main__':
    main()
